# emaktab_parser.py
"""
eMaktab tizimidan yuklab olingan rasmiy .xlsx va .xls elektron jurnallarni to'g'ridan-to'g'ri o'qish moduli.
"""

import os
import re
from typing import Dict, Any, List, Optional
import openpyxl
import xlrd
from cyrillic_translit import to_uzbek_latin

def clean_score(val: Any) -> Optional[int]:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return int(round(val)) if val > 0 else None
    s = str(val).strip().replace(",", ".")
    try:
        num = float(s)
        return int(round(num)) if num > 0 else None
    except (ValueError, TypeError):
        return None

def distribute_score_to_tasks(total_score: Optional[int], criteria: List[Dict[str, Any]]) -> List[Optional[int]]:
    if total_score is None or total_score <= 0 or not criteria:
        return [None] * len(criteria)
    
    max_total = sum(c.get("max_point", 10) for c in criteria)
    if max_total == 0:
        max_total = 50
        
    scores = []
    accumulated = 0
    
    for idx, c in enumerate(criteria):
        max_p = c.get("max_point", 10)
        # Proportionate allocation
        allocated = round((max_p / max_total) * total_score)
        allocated = min(allocated, max_p)
        scores.append(allocated)
        accumulated += allocated
        
    diff = total_score - accumulated
    # Adjust difference
    idx = 0
    while diff != 0 and idx < len(scores):
        max_p = criteria[idx].get("max_point", 10)
        if diff > 0 and scores[idx] < max_p:
            scores[idx] += 1
            diff -= 1
        elif diff < 0 and scores[idx] > 0:
            scores[idx] -= 1
            diff += 1
        idx = (idx + 1) % len(scores)
        
    return scores

def parse_emaktab_file(file_path: str, criteria: Optional[List[Dict[str, Any]]] = None) -> Dict[str, Any]:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".xls":
        return _parse_xls(file_path, criteria)
    else:
        return _parse_xlsx(file_path, criteria)

def _parse_xls(file_path: str, criteria: Optional[List[Dict[str, Any]]] = None) -> Dict[str, Any]:
    book = xlrd.open_workbook(file_path)
    sheet = book.sheet_by_index(0)

    class_name = ""
    subject = "Informatika"
    teacher_name = "Qobilov Obidjon"

    for r in range(min(6, sheet.nrows)):
        val = str(sheet.cell_value(r, 0))
        if "Sinf:" in val:
            m = re.search(r'Sinf:\s*([0-9]+-[A-ZА-ЯЁa-zа-яё]+)', val)
            if m:
                class_name = m.group(1)
        if "Fan:" in val:
            subject = val.replace("Fan:", "").strip()
        if "o‘qituvchining" in val.lower() or "o'qituvchining" in val.lower():
            t_raw = val.split(":")[-1].strip()
            if "Qobilov" in t_raw:
                teacher_name = "Qobilov Obidjon"
            elif t_raw:
                teacher_name = t_raw.split(",")[0].strip()

    # Find ALL BSB columns (there may be multiple, e.g. group 1, group 2 or 1-BSB / 2-BSB)
    bsb_cols = []
    max_score = 50
    for c in range(sheet.ncols):
        col_cells = [str(sheet.cell_value(r, c)) for r in range(min(25, sheet.nrows))]
        joined = " ".join(col_cells).upper()
        if any(k in joined for k in ["BSB", "БСБ", "CHSB", "ЧСБ", "СУММАТИВ", "SUMMATIV"]):
            bsb_cols.append(c)
            for cell_val in col_cells:
                m_score = re.search(r'\b(25|30|40|50)\b', cell_val)
                if m_score:
                    max_score = int(m_score.group(1))

    # If criteria is provided, criteria overrides max_score
    if criteria:
        max_score = sum(c.get("max_point", 10) for c in criteria)

    # Find student rows
    students = []
    start_r = -1
    for r in range(min(25, sheet.nrows)):
        val_0 = str(sheet.cell_value(r, 0)).lower()
        val_1 = str(sheet.cell_value(r, 1)).lower()
        if "ism" in val_1 or "фамилия" in val_1 or "ism" in val_0 or "фамилия" in val_0:
            start_r = r + 1
            break

    # If no BSB column was found by header, find columns with scores
    if not bsb_cols and start_r != -1:
        cand_counts = {}
        for c in range(2, sheet.ncols):
            cnt = sum(1 for r in range(start_r, sheet.nrows) if clean_score(sheet.cell_value(r, c)) is not None)
            if cnt > 0:
                cand_counts[c] = cnt
        if cand_counts:
            bsb_cols = sorted(cand_counts.keys(), key=lambda c: cand_counts[c], reverse=True)

    if start_r != -1 and bsb_cols:
        tr = 1
        for r in range(start_r, sheet.nrows):
            raw_name = str(sheet.cell_value(r, 1)).strip()
            if not raw_name or raw_name.isdigit() or len(raw_name) < 3:
                continue
            name = to_uzbek_latin(raw_name)
            
            # Check across all BSB columns: if first BSB column is empty, check the other BSB columns!
            tot = None
            for col in bsb_cols:
                score_val = sheet.cell_value(r, col)
                c_score = clean_score(score_val)
                if c_score is not None and c_score > 0:
                    tot = c_score
                    break

            task_scores = distribute_score_to_tasks(tot, criteria) if criteria else [None] * 5

            students.append({
                "tr": tr,
                "name": name,
                "total_score": tot,
                "task_scores": task_scores
            })
            tr += 1

    return {
        "school_info": {
            "class_name": to_uzbek_latin(class_name),
            "subject": to_uzbek_latin(subject),
            "quarter": 1,
            "bsb_number": 1,
            "max_score": max_score,
            "teacher_name": to_uzbek_latin(teacher_name),
            "head_teacher": "Mahmudova Dilfuzaxon"
        },
        "criteria": criteria or [],
        "students": students
    }

def _parse_xlsx(file_path: str, criteria: Optional[List[Dict[str, Any]]] = None) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active

    class_name = ""
    subject = "Informatika"
    teacher_name = "Qobilov Obidjon"

    for r in range(1, min(7, sheet.max_row + 1)):
        val = str(sheet.cell(r, 1).value or "")
        if "Sinf:" in val:
            m = re.search(r'Sinf:\s*([0-9]+-[A-ZА-ЯЁa-zа-яё]+)', val)
            if m:
                class_name = m.group(1)
        if "Fan:" in val:
            subject = val.replace("Fan:", "").strip()
        if "o‘qituvchining" in val.lower() or "o'qituvchining" in val.lower():
            t_raw = val.split(":")[-1].strip()
            if "Qobilov" in t_raw:
                teacher_name = "Qobilov Obidjon"
            elif t_raw:
                teacher_name = t_raw.split(",")[0].strip()

    # Find ALL BSB columns
    bsb_cols = []
    max_score = 50
    for c in range(1, sheet.max_column + 1):
        col_cells = [str(sheet.cell(r, c).value or "") for r in range(1, min(25, sheet.max_row + 1))]
        joined = " ".join(col_cells).upper()
        if any(k in joined for k in ["BSB", "БСБ", "CHSB", "ЧСБ", "СУММАТИВ", "SUMMATIV"]):
            bsb_cols.append(c)
            for cell_val in col_cells:
                m_score = re.search(r'\b(25|30|40|50)\b', cell_val)
                if m_score:
                    max_score = int(m_score.group(1))

    if criteria:
        max_score = sum(c.get("max_point", 10) for c in criteria)

    students = []
    start_r = -1
    for r in range(1, min(25, sheet.max_row + 1)):
        val_1 = str(sheet.cell(r, 1).value or "").lower()
        val_2 = str(sheet.cell(r, 2).value or "").lower()
        if "ism" in val_2 or "фамилия" in val_2 or "ism" in val_1 or "фамилия" in val_1:
            start_r = r + 1
            break

    name_col = 2 if start_r != -1 and "ism" in str(sheet.cell(start_r-1, 2).value or "").lower() else 1

    if not bsb_cols and start_r != -1:
        cand_counts = {}
        for c in range(name_col + 1, sheet.max_column + 1):
            cnt = sum(1 for r in range(start_r, sheet.max_row + 1) if clean_score(sheet.cell(r, c).value) is not None)
            if cnt > 0:
                cand_counts[c] = cnt
        if cand_counts:
            bsb_cols = sorted(cand_counts.keys(), key=lambda c: cand_counts[c], reverse=True)

    if start_r != -1 and bsb_cols:
        tr = 1
        for r in range(start_r, sheet.max_row + 1):
            raw_name = str(sheet.cell(r, name_col).value or "").strip()
            if not raw_name or raw_name.isdigit() or len(raw_name) < 3:
                continue
            name = to_uzbek_latin(raw_name)

            # Check across all BSB columns: if first column is empty, check the other BSB columns!
            tot = None
            for col in bsb_cols:
                score_val = sheet.cell(r, col).value
                c_score = clean_score(score_val)
                if c_score is not None and c_score > 0:
                    tot = c_score
                    break

            task_scores = distribute_score_to_tasks(tot, criteria) if criteria else [None] * 5

            students.append({
                "tr": tr,
                "name": name,
                "total_score": tot,
                "task_scores": task_scores
            })
            tr += 1

    return {
        "school_info": {
            "class_name": to_uzbek_latin(class_name),
            "subject": to_uzbek_latin(subject),
            "quarter": 1,
            "bsb_number": 1,
            "max_score": max_score,
            "teacher_name": to_uzbek_latin(teacher_name),
            "head_teacher": "Mahmudova Dilfuzaxon"
        },
        "criteria": criteria or [],
        "students": students
    }
