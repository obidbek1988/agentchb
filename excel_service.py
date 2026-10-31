from cyrillic_translit import to_uzbek_latin
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Dict, Any, List, Optional

thin_border = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000")
)

# Standard Fonts
font_title = Font(name="Times New Roman", size=11, bold=True)
font_header = Font(name="Times New Roman", size=8.5, bold=True)
font_data = Font(name="Times New Roman", size=10)
font_bold = Font(name="Times New Roman", size=10, bold=True)

# Colored Headers & Cells Styling
fill_hdr_fio = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")     # Soft ice blue
fill_hdr_tasks = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")   # Soft pastel blue
fill_hdr_score = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")   # Soft light green
fill_hdr_pct = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")     # Soft sky blue
fill_hdr_note = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")    # Soft amber yellow
fill_hdr_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

font_hdr_fio = Font(name="Times New Roman", size=9, bold=True, color="002060")
font_hdr_tasks = Font(name="Times New Roman", size=8.5, bold=True, color="1F4E79")
font_hdr_score = Font(name="Times New Roman", size=9, bold=True, color="276A3C")
font_hdr_pct = Font(name="Times New Roman", size=9, bold=True, color="1F4E79")
font_hdr_note = Font(name="Times New Roman", size=9, bold=True, color="7F6000")

# Data Row Colors
font_score_data = Font(name="Times New Roman", size=10, bold=True, color="1F4E79")
fill_score_data = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")

font_pct_data = Font(name="Times New Roman", size=10, bold=True, color="2E75B6")
fill_pct_data = PatternFill(start_color="F9FBFD", end_color="F9FBFD", fill_type="solid")

font_note_data = Font(name="Times New Roman", size=9.5, color="262626")

def generate_bsb_student_note(tot: Optional[float], max_score: Optional[int], task_scores: Optional[List[Any]], criteria: Optional[List[Dict[str, Any]]]) -> str:
    """Generates a detailed, natural explanation for BSB student score breakdown."""
    if tot is None or tot <= 0:
        return "Darsda qatnashmagan / baholanmagan"

    max_pts = max_score or 50
    if tot >= max_pts:
        return "Barcha topshiriqlar to‘liq va namunali bajarilgan (100%)."

    diff = int(round(max_pts - tot))

    if task_scores and criteria and len(task_scores) == len(criteria):
        full_tasks = []
        deducted_tasks = []
        for i, (ts, crit) in enumerate(zip(task_scores, criteria), start=1):
            mp = crit.get("max_point", 10)
            if ts is not None:
                if ts >= mp:
                    full_tasks.append(str(i))
                else:
                    deducted_tasks.append(f"{i}-topshiriqda (-{int(round(mp - ts))}b)")

        if full_tasks and deducted_tasks:
            full_str = ", ".join(full_tasks)
            ded_str = ", ".join(deducted_tasks)
            return f"{full_str}-topshiriq to‘liq; {ded_str} kamchilik mavjud."
        elif deducted_tasks:
            return f"{', '.join(deducted_tasks)} kamchiliklar mavjud."

    pct = (tot / max_pts) * 100
    if pct >= 85:
        return f"Topshiriqlar yaxshi bajarilgan, yakuniy qismda kichik xatolik (-{diff}b)."
    elif pct >= 70:
        return f"Nazariy qism to‘g‘ri, amaliy mashq va loyihada kamchiliklar bor (-{diff}b)."
    elif pct >= 50:
        return f"Topshiriqlar qisman bajarilgan, amaliyotda xatoliklar mavjud (-{diff}b)."
    else:
        return f"Topshiriqlar chala qolgan, nazariya va amaliyotda xatolar ko‘p (-{diff}b)."

def add_monitoring_summary_block(
    ws,
    start_row: int,
    end_row: int,
    max_score: int,
    score_col_letter: str,
    label_merge_start: str,
    label_merge_end: str,
    score_col_idx: int,
    total_cols: int,
    lang: str = "uz"
) -> int:
    """
    Adds official colored monitoring summary table (A'lo, Yaxshi, Qoniqarli, Qoniqarsiz,
    Qatnashmadi, Jami o'quvchilar soni, O'zlashtirish %, Bilim sifati %) with live formulas.
    Returns the next row index for signatures.
    """
    if max_score == 50:
        alo_min, alo_max = 43, 50
        yax_min, yax_max = 33, 42
        qon_min, qon_max = 15, 32
        qonsiz_max = 14
    elif max_score == 25:
        alo_min, alo_max = 22, 25
        yax_min, yax_max = 17, 21
        qon_min, qon_max = 8, 16
        qonsiz_max = 7
    elif max_score == 30:
        alo_min, alo_max = 26, 30
        yax_min, yax_max = 20, 25
        qon_min, qon_max = 9, 19
        qonsiz_max = 8
    elif max_score == 40:
        alo_min, alo_max = 35, 40
        yax_min, yax_max = 27, 34
        qon_min, qon_max = 13, 26
        qonsiz_max = 12
    else:
        alo_min = int(round(max_score * 0.86))
        yax_min = int(round(max_score * 0.66))
        qon_min = int(round(max_score * 0.30))
        alo_max = max_score
        yax_max = alo_min - 1
        qon_max = yax_min - 1
        qonsiz_max = qon_min - 1

    r = end_row + 1
    sc = score_col_letter

    lbl_alo = f"A‘lo  ({alo_min}–{alo_max})" if lang != "ru" else f"Отлично  ({alo_min}–{alo_max})"
    lbl_yax = f"Yaxshi  ({yax_min}–{yax_max})" if lang != "ru" else f"Хорошо  ({yax_min}–{yax_max})"
    lbl_qon = f"Qoniqarli  ({qon_min}–{qon_max})" if lang != "ru" else f"Удовлетворительно  ({qon_min}–{qon_max})"
    lbl_qonsiz = f"Qoniqarsiz  (0–{qonsiz_max})" if lang != "ru" else f"Неудовлетворительно  (0–{qonsiz_max})"
    lbl_qatn = "Qatnashmadi" if lang != "ru" else "Не участвовал"
    lbl_jami = "Jami o‘quvchilar soni" if lang != "ru" else "Всего учащихся"
    lbl_ozl = "O‘zlashtirish" if lang != "ru" else "Успеваемость"
    lbl_bil = "Bilim sifati" if lang != "ru" else "Качество знаний"

    summary_rows = [
        (r,   lbl_alo,    f"=COUNTIF({sc}{start_row}:{sc}{end_row}, \">={alo_min}\")", "C00000", "FCE4D6", True, None),
        (r+1, lbl_yax,    f"=COUNTIF({sc}{start_row}:{sc}{end_row}, \">={yax_min}\")-COUNTIF({sc}{start_row}:{sc}{end_row}, \">={alo_min}\")", "385723", "E2EFDA", True, None),
        (r+2, lbl_qon,    f"=COUNTIF({sc}{start_row}:{sc}{end_row}, \">={qon_min}\")-COUNTIF({sc}{start_row}:{sc}{end_row}, \">={yax_min}\")", "1F4E79", "D9E1F2", True, None),
        (r+3, lbl_qonsiz, f"=COUNTIF({sc}{start_row}:{sc}{end_row}, \"<{qon_min}\")-COUNTBLANK({sc}{start_row}:{sc}{end_row})", "833C0C", "F2F2F2", True, None),
        (r+4, lbl_qatn,   f"=COUNTBLANK({sc}{start_row}:{sc}{end_row})", "595959", "F2F2F2", False, None),
        (r+5, lbl_jami,   f"=COUNT({sc}{start_row}:{sc}{end_row})+COUNTBLANK({sc}{start_row}:{sc}{end_row})", "000000", "D9D9D9", True, None),
        (r+6, lbl_ozl,    f"=IF({sc}{r+5}>0, ({sc}{r}+{sc}{r+1}+{sc}{r+2})/{sc}{r+5}, 0)", "203764", "DDEBF7", True, "0.0%"),
        (r+7, lbl_bil,    f"=IF({sc}{r+5}>0, ({sc}{r}+{sc}{r+1})/{sc}{r+5}, 0)", "C00000", "FCE4D6", True, "0.0%")
    ]

    for row_idx, label, formula, font_color, fill_color, is_bold, num_fmt in summary_rows:
        ws.row_dimensions[row_idx].height = 20
        ws.merge_cells(f"{label_merge_start}{row_idx}:{label_merge_end}{row_idx}")
        lbl_cell = ws.cell(row=row_idx, column=2, value=label)
        lbl_cell.font = Font(name="Times New Roman", size=10, bold=is_bold, color=font_color)
        lbl_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        val_cell = ws.cell(row=row_idx, column=score_col_idx, value=formula)
        val_cell.font = Font(name="Times New Roman", size=10, bold=is_bold, color=font_color)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        if num_fmt:
            val_cell.number_format = num_fmt

        fill_obj = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid") if fill_color else None

        for c_idx in range(1, total_cols + 1):
            c = ws.cell(row=row_idx, column=c_idx)
            c.border = thin_border
            if fill_obj and (2 <= c_idx <= score_col_idx):
                c.fill = fill_obj

    return r + 8

def build_bsb_excel(data: Dict[str, Any], output_path: str) -> str:
    """BSB standard 50/25/30 ball report (A4 Portrait, exact official template with colored monitoring summary)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.views.sheetView[0].showGridLines = True

    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

    school_info = data.get("school_info", {})
    criteria = data.get("criteria", [])
    students = data.get("students", [])

    class_name = to_uzbek_latin(school_info.get("class_name", ""))
    subject = school_info.get("subject", "Informatika")
    quarter = school_info.get("quarter", 1)
    bsb_number = school_info.get("bsb_number", 1)
    max_score = school_info.get("max_score", 50)
    teacher_name = to_uzbek_latin(school_info.get("teacher_name", "Qobilov Obidjon"))
    head_teacher = to_uzbek_latin(school_info.get("head_teacher", "Mahmudova Dilfuzaxon"))

    title_text = (
        "BSB tahlil Bayonnomasi\n"
        f"FARG‘ONA viloyati FARG‘ONA SHAHRI 26-umumta’lim maktabi\n"
        f"{class_name} sinf o‘quvchilarining {quarter}-chorak {subject.lower()} fanidan {bsb_number}-BSB ({max_score}ball)natijalari"
    )

    ws.merge_cells("A1:I1")
    c_a1 = ws["A1"]
    c_a1.value = title_text
    c_a1.font = font_title
    c_a1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45

    # Headers with colored background fills
    c_tr = ws.cell(row=2, column=1, value="t/r")
    c_tr.font = font_hdr_fio
    c_tr.fill = fill_hdr_fio

    c_fio = ws.cell(row=2, column=2, value="O‘quvchilarning ism-sharifi")
    c_fio.font = font_hdr_fio
    c_fio.fill = fill_hdr_fio

    for idx, crit in enumerate(criteria[:5], start=3):
        txt = crit.get("skill_text", f"{idx-2}-topshiriq")
        c = ws.cell(row=2, column=idx, value=txt)
        c.font = font_hdr_tasks
        c.fill = fill_hdr_tasks
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    c_sc = ws.cell(row=2, column=8, value=f"Jami ball\n{max_score}")
    c_sc.font = font_hdr_score
    c_sc.fill = fill_hdr_score

    c_pc = ws.cell(row=2, column=9, value="Foizi\n100%")
    c_pc.font = font_hdr_pct
    c_pc.fill = fill_hdr_pct

    c_nt = ws.cell(row=2, column=10, value="Tahlil / Qisqa izoh\n(Xatolar va to‘g‘rilar)")
    c_nt.font = font_hdr_note
    c_nt.fill = fill_hdr_note

    for col in [1, 2, 8, 9, 10]:
        cell = ws.cell(row=2, column=col)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[2].height = 145

    start_row = 3
    for idx, st in enumerate(students, start=1):
        curr_row = start_row + idx - 1
        name = to_uzbek_latin(st.get("name", ""))
        tot = st.get("total_score")
        task_scores = st.get("task_scores")

        if tot is not None and tot > 0:
            if not task_scores or all(s is None for s in task_scores) or sum(s or 0 for s in task_scores) != tot:
                from emaktab_parser import distribute_score_to_tasks
                task_scores = distribute_score_to_tasks(tot, criteria)
        else:
            task_scores = [None] * len(criteria)

        c_a = ws.cell(row=curr_row, column=1, value=idx)
        c_a.alignment = Alignment(horizontal="center", vertical="center")
        c_a.border = thin_border
        c_a.font = font_data

        c_b = ws.cell(row=curr_row, column=2, value=name)
        c_b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c_b.border = thin_border
        c_b.font = font_data

        for c_idx, score in enumerate(task_scores[:5], start=3):
            val = score if score is not None else ""
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            cell.font = font_data

        c_h = ws.cell(row=curr_row, column=8)
        if tot is not None:
            c_h.value = f"=SUM(C{curr_row}:G{curr_row})"
        else:
            c_h.value = ""
        c_h.alignment = Alignment(horizontal="center", vertical="center")
        c_h.border = thin_border
        c_h.font = font_score_data
        c_h.fill = fill_score_data

        c_i = ws.cell(row=curr_row, column=9)
        if tot is not None:
            if max_score == 50:
                c_i.value = f"=IF(H{curr_row}=\"\",\"\",H{curr_row}*2)"
            elif max_score == 25:
                c_i.value = f"=IF(H{curr_row}=\"\",\"\",H{curr_row}*4)"
            elif max_score == 30:
                c_i.value = f"=IF(H{curr_row}=\"\",\"\",ROUND(H{curr_row}/30*100,0))"
            else:
                c_i.value = f"=IF(H{curr_row}=\"\",\"\",ROUND(H{curr_row}/{max_score}*100,0))"
        else:
            c_i.value = ""
        c_i.alignment = Alignment(horizontal="center", vertical="center")
        c_i.border = thin_border
        c_i.font = font_pct_data
        c_i.fill = fill_pct_data

        note = st.get("note")
        if not note or str(note).strip() in ["", "Tekshirildi", "None"]:
            note = generate_bsb_student_note(tot, max_score, task_scores, criteria)
        c_j = ws.cell(row=curr_row, column=10, value=to_uzbek_latin(str(note)))
        c_j.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        c_j.border = thin_border
        c_j.font = font_note_data

        ws.row_dimensions[curr_row].height = 20

    end_row = start_row + len(students) - 1

    # Insert official colored monitoring summary block (O'quvchilar tugashi va o'qituvchi tepasiga)
    next_avail_row = add_monitoring_summary_block(
        ws,
        start_row=start_row,
        end_row=end_row,
        max_score=max_score,
        score_col_letter="H",
        label_merge_start="B",
        label_merge_end="G",
        score_col_idx=8,
        total_cols=10,
        lang="uz"
    )

    sig_row = next_avail_row + 1
    ws.cell(row=sig_row, column=1, value=f"Fan o‘qituvchisi:____________________  {teacher_name}").font = font_bold
    ws.cell(row=sig_row+1, column=1, value=f"Metod birlashma rahbari:_______________ {head_teacher}").font = font_bold

    # Print area strictly covers through Foizi column including summary and signatures
    ws.print_area = f"A1:I{sig_row+1}"

    ws.column_dimensions["A"].width = 4.5
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 13
    ws.column_dimensions["F"].width = 13
    ws.column_dimensions["G"].width = 13
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 9
    ws.column_dimensions["J"].width = 65

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path

def build_chsb_excel(data: Dict[str, Any], output_path: str) -> str:
    """CHSB/SOCH 40 ball report: (B/З) Bilish/Znanie, (Q/П) Qo'llash/Primenenie, (M/Р) Mulohaza/Rassujdenie template with colored monitoring summary"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.views.sheetView[0].showGridLines = True

    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

    school_info = data.get("school_info", {})
    students = data.get("students", [])

    lang = school_info.get("lang", "uz")
    if "рус" in str(school_info.get("class_name", "")).lower() or "rus" in str(school_info.get("class_name", "")).lower() or lang == "ru":
        lang = "ru"

    quarter = school_info.get("quarter", 1)
    chsb_number = school_info.get("chsb_number", 1)
    max_score = school_info.get("max_score", 40)

    if lang == "ru":
        class_name = school_info.get("class_name", "9-класс")
        subject = school_info.get("subject", "Информатика")
        teacher_name = school_info.get("teacher_name", "Кабилов Обиджон")
        head_teacher = school_info.get("head_teacher", "Махмудова Дильфузахон")
        title_text = (
            "Протокол анализа результатов СОЧ\n"
            f"ФЕРГАНСКАЯ область город ФЕРГАНА общеобразовательная школа №26\n"
            f"Результаты СОЧ-1 ({max_score} баллов) по предмету «{subject}» учащихся {class_name} за {quarter}-ю четверть"
        )
        chsb_headers = [
            ("A2", "№", fill_hdr_fio, font_hdr_fio),
            ("B2", "Ф.И.О. учащихся", fill_hdr_fio, font_hdr_fio),
            ("C2", "(З) – Знание,\nза каждый верный ответ – 2 балла,\nвсего 6 баллов;", fill_hdr_tasks, font_hdr_tasks),
            ("D2", "(П) – Применение,\nза каждый верный ответ – 2,4 балла,\nвсего 24 балла;", fill_hdr_tasks, font_hdr_tasks),
            ("E2", "(Р) – Рассуждение,\nза каждый верный ответ – 5 баллов,\nвсего 10 баллов;", fill_hdr_tasks, font_hdr_tasks),
            ("F2", f"Итоговый балл\n{max_score}", fill_hdr_score, font_hdr_score),
            ("G2", "Процент\n100%", fill_hdr_pct, font_hdr_pct),
            ("H2", "Анализ / Краткий комментарий\n(Ошибки и верные ответы)", fill_hdr_note, font_hdr_note),
            ("I2", "Ответы, отмеченные учащимся\n(1–15 ответы теста)", fill_hdr_gray, font_header)
        ]
    else:
        class_name = to_uzbek_latin(school_info.get("class_name", ""))
        subject = school_info.get("subject", "Informatika")
        teacher_name = to_uzbek_latin(school_info.get("teacher_name", "Qobilov Obidjon"))
        head_teacher = to_uzbek_latin(school_info.get("head_teacher", "Mahmudova Dilfuzaxon"))
        title_text = (
            "CHSB test tahlil Bayonnomasi\n"
            f"FARG‘ONA viloyati FARG‘ONA SHAHRI 26-umumta’lim maktabi\n"
            f"{class_name} sinf o‘quvchilarining {quarter}-chorak {subject.lower()} fanidan {chsb_number}-CHSB ({max_score} ball) test natijalari"
        )
        chsb_headers = [
            ("A2", "t/r", fill_hdr_fio, font_hdr_fio),
            ("B2", "O‘quvchilarning ism-sharifi", fill_hdr_fio, font_hdr_fio),
            ("C2", "(B) – Bilish,\nhar bir to‘g‘ri javob uchun – 2 balldan,\njami 6 ball;", fill_hdr_tasks, font_hdr_tasks),
            ("D2", "(Q) – Qo‘llash,\nhar bir to‘g‘ri javob uchun – 2,4 balldan,\njami 24 ball;", fill_hdr_tasks, font_hdr_tasks),
            ("E2", "(M) – Mulohaza,\nhar bir to‘g‘ri javob uchun – 5 balldan,\njami 10 ball;", fill_hdr_tasks, font_hdr_tasks),
            ("F2", f"Jami ball\n{max_score}", fill_hdr_score, font_hdr_score),
            ("G2", "Foizi\n100%", fill_hdr_pct, font_hdr_pct),
            ("H2", "Tahlil / Qisqa izoh\n(Xatolar va to‘g‘rilar)", fill_hdr_note, font_hdr_note),
            ("I2", "O‘quvchi belgilagan kalitlar\n(1–15 test javoblari)", fill_hdr_gray, font_header)
        ]

    ws.merge_cells("A1:G1")
    c_a1 = ws["A1"]
    c_a1.value = title_text
    c_a1.font = font_title
    c_a1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45

    for coord, txt, fill_style, font_style in chsb_headers:
        cell = ws[coord]
        cell.value = txt
        cell.font = font_style
        cell.fill = fill_style
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[2].height = 95

    start_row = 3
    for idx, st in enumerate(students, start=1):
        curr_row = start_row + idx - 1
        name = st.get("name", "") if lang == "ru" else to_uzbek_latin(st.get("name", ""))
        bqm = st.get("task_scores", [None, None, None])
        tot = st.get("total_score")

        c_a = ws.cell(row=curr_row, column=1, value=idx)
        c_a.alignment = Alignment(horizontal="center", vertical="center")
        c_a.border = thin_border
        c_a.font = font_data

        c_b = ws.cell(row=curr_row, column=2, value=name)
        c_b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c_b.border = thin_border
        c_b.font = font_data

        for c_idx, score in enumerate(bqm[:3], start=3):
            val = score if score is not None else ""
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            cell.font = font_data

        c_f = ws.cell(row=curr_row, column=6)
        if tot is not None:
            c_f.value = f"=SUM(C{curr_row}:E{curr_row})"
        else:
            c_f.value = ""
        c_f.alignment = Alignment(horizontal="center", vertical="center")
        c_f.border = thin_border
        c_f.font = font_score_data
        c_f.fill = fill_score_data

        c_g = ws.cell(row=curr_row, column=7)
        if tot is not None:
            c_g.value = f"=IF(F{curr_row}=\"\",\"\",ROUND(F{curr_row}/40*100,0))"
        else:
            c_g.value = ""
        c_g.alignment = Alignment(horizontal="center", vertical="center")
        c_g.border = thin_border
        c_g.font = font_pct_data
        c_g.fill = fill_pct_data

        note = st.get("note", "") if lang == "ru" else to_uzbek_latin(st.get("note", ""))
        c_h = ws.cell(row=curr_row, column=8, value=note)
        c_h.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        c_h.border = thin_border
        c_h.font = font_note_data

        keys_val = st.get("keys_str") or st.get("student_keys", "")
        c_i = ws.cell(row=curr_row, column=9, value=keys_val)
        c_i.alignment = Alignment(horizontal="left", vertical="center")
        c_i.border = thin_border
        c_i.font = font_data

        ws.row_dimensions[curr_row].height = 19

    end_row = start_row + len(students) - 1

    # Insert official colored monitoring summary block (O'quvchilar tugashi va o'qituvchi tepasiga)
    next_avail_row = add_monitoring_summary_block(
        ws,
        start_row=start_row,
        end_row=end_row,
        max_score=max_score,
        score_col_letter="F",
        label_merge_start="B",
        label_merge_end="E",
        score_col_idx=6,
        total_cols=9,
        lang=lang
    )

    sig_row = next_avail_row + 1
    if lang == "ru":
        ws.cell(row=sig_row, column=1, value=f"Учитель предмета:____________________  {teacher_name}").font = font_bold
        ws.cell(row=sig_row+1, column=1, value=f"Заместитель директора по УВР:_______________ {head_teacher}").font = font_bold
    else:
        ws.cell(row=sig_row, column=1, value=f"Fan o‘qituvchisi:____________________  {teacher_name}").font = font_bold
        ws.cell(row=sig_row+1, column=1, value=f"Metod birlashma rahbari:_______________ {head_teacher}").font = font_bold

    # Print area strictly covers through Foizi column including summary and signatures
    ws.print_area = f"A1:G{sig_row+1}"

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 17
    ws.column_dimensions["D"].width = 19
    ws.column_dimensions["E"].width = 17
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 9
    ws.column_dimensions["H"].width = 60
    ws.column_dimensions["I"].width = 48

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path
